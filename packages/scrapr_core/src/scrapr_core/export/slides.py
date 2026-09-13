"""The PowerPoint renderer (`REQ-EXP-002`, `DEC-21`).

A pure function from an `ExportDocument` and a `Theme` to the bytes of a `.pptx`.

**Content is distributed by the report's structure** (`AC-1`): a cover, What's
Changed when there is one, each section's claims across as many slides as they
need, a slide per chart, and the sources. Nothing is truncated; a section with
many claims becomes several slides, each saying which of how many it is.

**Charts are native PowerPoint charts** (`AC-2`), built from the spec's points,
so they are real, editable slide content rather than pictures of charts.
Tables are native tables.

**The file must open without a repair prompt** (`AC-3`). Every string has
already been stripped of characters XML cannot hold (`document.clean_text`),
only `http`/`https` links are attached, and the deck is built from the default
template's blank layout, which is the path python-pptx is best exercised on.
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from typing import Final

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Emu, Inches, Pt

from scrapr_core.db.enums import VizKind
from scrapr_core.export.document import (
    ExportChart,
    ExportClaim,
    ExportDocument,
    ExportReference,
    format_date,
    format_value,
)
from scrapr_core.export.themes import RULES, Cover, Rule, Theme

__all__ = ["render_pptx"]

SLIDE_WIDTH: Final = Inches(13.333)
SLIDE_HEIGHT: Final = Inches(7.5)
MARGIN: Final = Inches(0.7)
CONTENT_TOP: Final = Inches(1.55)
CONTENT_BOTTOM: Final = Inches(6.7)
CLAIM_BUDGET: Final = 6
"""How much claim a slide holds, in units of roughly one short claim."""
SOURCES_PER_SLIDE: Final = 9

_DASH: Final = {
    Rule.SOLID: MSO_LINE_DASH_STYLE.SOLID,
    Rule.DASHED: MSO_LINE_DASH_STYLE.DASH,
    Rule.DOTTED: MSO_LINE_DASH_STYLE.ROUND_DOT,
}


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#").upper())


class _Deck:
    """A presentation being built in one theme."""

    def __init__(self, document: ExportDocument, theme: Theme) -> None:
        self.document = document
        self.theme = theme
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_WIDTH
        self.prs.slide_height = SLIDE_HEIGHT
        self._blank = self.prs.slide_layouts[6]
        properties = self.prs.core_properties
        properties.title = document.objective[:250]
        properties.subject = document.version_line
        properties.author = "ScrapR"

    # -- primitives ---------------------------------------------------------

    def slide(self, title: str | None = None):  # type: ignore[no-untyped-def]
        slide = self.prs.slides.add_slide(self._blank)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = _rgb(self.theme.ground)
        if self.theme.cover is Cover.BAND:
            band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, Inches(0.12))
            self.fill(band, self.theme.accent)
        if title is not None:
            self.text(
                slide, title, MARGIN, Inches(0.55), SLIDE_WIDTH - 2 * MARGIN, Inches(0.8),
                size=26 * self.theme.title_scale, bold=self.theme.heading_bold,
                font=self.theme.pptx_heading_font,
            )
        self.footer(slide)
        return slide

    def fill(self, shape, colour: str) -> None:  # type: ignore[no-untyped-def]
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(colour)
        shape.line.fill.background()

    def text(  # type: ignore[no-untyped-def]
        self, slide, value: str, left: int, top: int, width: int, height: int, *,
        size: float = 14, bold: bool = False, colour: str | None = None, font: str | None = None,
    ):
        box = slide.shapes.add_textbox(left, top, width, height)
        frame = box.text_frame
        frame.word_wrap = True
        run = frame.paragraphs[0].add_run()
        run.text = value
        self.style(run, size=size, bold=bold, colour=colour, font=font)
        return box

    def style(  # type: ignore[no-untyped-def]
        self, run, *, size: float, bold: bool = False, colour: str | None = None, font: str | None = None,
    ) -> None:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = font or self.theme.pptx_body_font
        run.font.color.rgb = _rgb(colour or self.theme.text)

    def footer(self, slide) -> None:  # type: ignore[no-untyped-def]
        """`REQ-EXP-006 AC-2`: every slide says which version it came from."""
        number = len(self.prs.slides)
        self.text(
            slide, f"{self.document.version_line}  ·  {number}", MARGIN, Inches(6.95),
            SLIDE_WIDTH - 2 * MARGIN, Inches(0.35), size=9, colour=self.theme.muted,
        )

    # -- slides -------------------------------------------------------------

    def cover(self) -> None:
        theme, document = self.theme, self.document
        slide = self.prs.slides.add_slide(self._blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(theme.ground)

        title_colour = theme.text
        if theme.cover is Cover.ACCENT_BLOCK:
            block = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, Inches(4.6))
            self.fill(block, theme.accent)
            title_colour = theme.accent_on_accent
        elif theme.cover is Cover.BAND:
            band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, Inches(0.6))
            self.fill(band, theme.accent)

        top = Inches(1.6) if theme.cover is not Cover.PLAIN else Inches(2.6)
        # Smaller type for a longer objective, so a long brief stays on the
        # cover rather than running over the accent block and the version line.
        length = len(document.objective)
        size = 36.0 if length <= 90 else 28.0 if length <= 180 else 20.0 if length <= 400 else 14.0
        self.text(
            slide, document.objective, MARGIN, top, SLIDE_WIDTH - 2 * MARGIN, Inches(2.2),
            size=size * theme.title_scale, bold=theme.heading_bold, colour=title_colour,
            font=theme.pptx_heading_font,
        )
        meta_colour = theme.accent_on_accent if theme.cover is Cover.ACCENT_BLOCK else theme.muted
        lines = [document.version_line]
        if document.subject:
            lines.insert(0, f"Subject: {document.subject}")
        self.text(slide, "\n".join(lines), MARGIN, Inches(3.8), SLIDE_WIDTH - 2 * MARGIN,
                  Inches(0.7), size=14, colour=meta_colour)

        if theme.cover is Cover.ACCENT_RULE:
            rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(4.7), Inches(1.8), Inches(0.06))
            self.fill(rule, theme.accent)
        if theme.cover is Cover.FIGURES_STRIP:
            self.figures_strip(slide)
        if document.is_partial:
            self.text(
                slide,
                "Parts of this question could not be answered. What is missing is marked in "
                "the report, and nothing has been filled in with guesswork.",
                MARGIN, Inches(5.8), SLIDE_WIDTH - 2 * MARGIN, Inches(0.7), size=12,
                colour=theme.text,
            )
        self.footer(slide)

    def figures_strip(self, slide) -> None:  # type: ignore[no-untyped-def]
        cells: list[tuple[str, str]] = []
        for section in self.document.sections:
            for chart in section.charts:
                for series in chart.series:
                    if series.points and len(cells) < 3:
                        last_label, last_value = series.points[-1]
                        cells.append(
                            (format_value(last_value, chart.unit), f"{chart.title or series.name}, {last_label}")
                        )
        if not cells:
            claims = sum(len(section.claims) for section in self.document.sections)
            cells = [(str(claims), "findings"), (str(len(self.document.references)), "sources"),
                     (str(len(self.document.sections)), "sections")]
        width = (SLIDE_WIDTH - 2 * MARGIN) // len(cells)
        for index, (value, caption) in enumerate(cells):
            left = MARGIN + width * index
            self.text(slide, value, left, Inches(4.8), width, Inches(0.7), size=28, bold=True,
                      colour=self.theme.accent, font=self.theme.pptx_heading_font)
            self.text(slide, caption, left, Inches(5.45), width, Inches(0.5), size=11,
                      colour=self.theme.muted)

    def changes(self) -> None:
        document = self.document
        if document.changes_headline is None:
            return
        chunks = _chunk(list(document.changes), 5) or [[]]
        for index, chunk in enumerate(chunks):
            # A typographic apostrophe, as the workspace heading has it.
            title = "What’s changed" + (f" ({index + 1} of {len(chunks)})" if len(chunks) > 1 else "")  # noqa: RUF001
            slide = self.slide(title)
            box = slide.shapes.add_textbox(MARGIN, CONTENT_TOP, SLIDE_WIDTH - 2 * MARGIN, CONTENT_BOTTOM - CONTENT_TOP)
            frame = box.text_frame
            frame.word_wrap = True
            first = frame.paragraphs[0].add_run()
            first.text = document.changes_headline
            self.style(first, size=18)
            for change in chunk:
                label = frame.add_paragraph()
                label.space_before = Pt(10)
                label_run = label.add_run()
                label_run.text = change.label
                self.style(label_run, size=11, bold=True, colour=self.theme.muted)
                body = frame.add_paragraph()
                run = body.add_run()
                run.text = change.summary
                self.style(run, size=13)

    def section(self, title: str, claims: Sequence[ExportClaim], charts: Sequence[ExportChart]) -> None:
        pages = _paginate_claims(claims)
        for index, page in enumerate(pages):
            suffix = f" ({index + 1} of {len(pages)})" if len(pages) > 1 else ""
            slide = self.slide(title + suffix)
            top = int(CONTENT_TOP)
            available = CONTENT_BOTTOM - CONTENT_TOP
            unit = int(available) // CLAIM_BUDGET
            for claim in page:
                height = int(unit * _weight(claim))
                self.claim(slide, claim, top, height)
                top += height
        for chart in charts:
            self.chart(chart)

    def claim(self, slide, claim: ExportClaim, top: int, height: int) -> None:  # type: ignore[no-untyped-def]
        theme = self.theme
        rule = RULES[claim.claim_type]
        if rule is not Rule.NONE:
            line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, MARGIN, top + Inches(0.08),
                                              MARGIN, top + height - Inches(0.12))
            line.line.color.rgb = _rgb(theme.text)
            line.line.width = Pt(2)
            line.line.dash_style = _DASH[rule]

        box = slide.shapes.add_textbox(MARGIN + Inches(0.2), top, SLIDE_WIDTH - 2 * MARGIN - Inches(0.2), height)
        frame = box.text_frame
        frame.word_wrap = True

        label = claim.type_word + (f" · {claim.confidence}" if claim.confidence else "")
        head = frame.paragraphs[0].add_run()
        head.text = label
        self.style(head, size=10, bold=True, colour=theme.muted)

        body = frame.add_paragraph()
        text_run = body.add_run()
        text_run.text = claim.text
        self.style(text_run, size=14)
        if claim.references:
            refs = body.add_run()
            refs.text = " " + "".join(f"[{number}]" for number in claim.references)
            self.style(refs, size=10, colour=theme.accent if theme.accent != theme.text else theme.muted)

        if claim.reporting_period:
            self.meta(frame, f"Covers the period ending {claim.reporting_period}")
        if claim.assumptions:
            self.meta(frame, "Assumptions: " + "; ".join(claim.assumptions))
        for conflict in claim.conflicts:
            sides = "; ".join(
                f"{side.value} from {side.source_name}"
                + (" (your document)" if side.from_your_document else "")
                + (f" [{side.reference}]" if side.reference else "")
                for side in conflict.sides
            )
            heading = "Sources disagree" if not conflict.resolved else "Sources differ, explained"
            self.meta(frame, f"{heading}: {sides}. {conflict.explanation}")

    def meta(self, frame, value: str) -> None:  # type: ignore[no-untyped-def]
        paragraph = frame.add_paragraph()
        run = paragraph.add_run()
        run.text = value
        self.style(run, size=10, colour=self.theme.muted)

    def chart(self, chart: ExportChart) -> None:
        theme = self.theme
        slide = self.slide(chart.title or "Chart")
        left, top = MARGIN, CONTENT_TOP
        width, height = SLIDE_WIDTH - 2 * MARGIN, Inches(4.6)

        if chart.kind in {VizKind.TABLE, VizKind.MATRIX}:
            rows = [(series.name, label, format_value(value, chart.unit))
                    for series in chart.series for label, value in series.points]
            shape = slide.shapes.add_table(len(rows) + 1, 3, left, top, width, Emu(min(height, Inches(0.4) * (len(rows) + 1))))
            table = shape.table
            for column, heading in enumerate(("Series", "Label", "Value")):
                table.cell(0, column).text = heading
            for row_index, row in enumerate(rows, start=1):
                for column, cell_text in enumerate(row):
                    table.cell(row_index, column).text = cell_text
            for row_cells in table.rows:
                for cell in row_cells.cells:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _rgb(theme.surface)
                    for paragraph in cell.text_frame.paragraphs:
                        for run in paragraph.runs:
                            self.style(run, size=12)
        elif chart.kind is VizKind.METRIC:
            index = 0
            for series in chart.series:
                for label, value in series.points:
                    self.text(slide, format_value(value, chart.unit), left, top + Inches(1.1) * index,
                              width, Inches(0.8), size=36, bold=True, colour=theme.accent,
                              font=theme.pptx_heading_font)
                    self.text(slide, f"{series.name} {label}".strip(), left, top + Inches(0.75) + Inches(1.1) * index,
                              width, Inches(0.4), size=12, colour=theme.muted)
                    index += 1
        else:
            data = CategoryChartData()
            labels = list(chart.labels)
            data.categories = labels
            for series in chart.series:
                by_label = dict(series.points)
                data.add_series(series.name or chart.title or "Series",
                                [float(by_label[label]) if label in by_label else None for label in labels])
            kind = XL_CHART_TYPE.LINE_MARKERS if chart.kind is VizKind.LINE else XL_CHART_TYPE.COLUMN_CLUSTERED
            frame = slide.shapes.add_chart(kind, left, top, width, height, data)
            plotted = frame.chart
            plotted.has_legend = len(chart.series) > 1
            if plotted.has_legend:
                plotted.legend.position = XL_LEGEND_POSITION.BOTTOM
                plotted.legend.include_in_layout = False
                plotted.legend.font.color.rgb = _rgb(theme.text)
            plotted.font.size = Pt(11)
            plotted.font.name = theme.pptx_body_font
            plotted.font.color.rgb = _rgb(theme.text)
            for index, series_plot in enumerate(plotted.plots[0].series):
                colour = _rgb(theme.accent if index == 0 else theme.neutral_series)
                if chart.kind is VizKind.LINE:
                    series_plot.format.line.color.rgb = colour
                    series_plot.format.line.width = Pt(2.25)
                    if index:
                        series_plot.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
                    series_plot.marker.format.fill.solid()
                    series_plot.marker.format.fill.fore_color.rgb = colour
                else:
                    series_plot.format.fill.solid()
                    series_plot.format.fill.fore_color.rgb = colour
            if chart.unit:
                self.text(slide, f"Values in {chart.unit}", left, top + height + Inches(0.05), width,
                          Inches(0.35), size=10, colour=theme.muted)

        if chart.references:
            self.text(slide, "Sources: " + "".join(f"[{n}]" for n in chart.references), left,
                      Inches(6.45), width, Inches(0.35), size=10, colour=theme.muted)

    def sources(self, references: Sequence[ExportReference]) -> None:
        chunks = _chunk(list(references), SOURCES_PER_SLIDE)
        for index, chunk in enumerate(chunks):
            title = "Sources" + (f" ({index + 1} of {len(chunks)})" if len(chunks) > 1 else "")
            slide = self.slide(title)
            box = slide.shapes.add_textbox(MARGIN, CONTENT_TOP, SLIDE_WIDTH - 2 * MARGIN, CONTENT_BOTTOM - CONTENT_TOP)
            frame = box.text_frame
            frame.word_wrap = True
            for position, reference in enumerate(chunk):
                paragraph = frame.paragraphs[0] if position == 0 else frame.add_paragraph()
                paragraph.space_after = Pt(6)
                marker = paragraph.add_run()
                marker.text = f"[{reference.number}] "
                self.style(marker, size=12, bold=True)
                name = paragraph.add_run()
                name.text = reference.name
                self.style(name, size=12, colour=self.theme.text)
                if reference.url:
                    name.hyperlink.address = reference.url
                details = [reference.tier, f"read {format_date(reference.retrieved_at)}"]
                if reference.publisher:
                    details.insert(0, reference.publisher)
                if reference.from_your_document:
                    details.insert(0, "your document")
                if reference.document_removed:
                    details.append("file since deleted")
                rest = paragraph.add_run()
                rest.text = ". " + ", ".join(details) + "."
                self.style(rest, size=11, colour=self.theme.muted)

    def bytes(self) -> bytes:
        buffer = io.BytesIO()
        self.prs.save(buffer)
        return buffer.getvalue()


def _weight(claim: ExportClaim) -> int:
    """Roughly how many short-claim slots a claim needs on a slide."""
    extra = len(claim.conflicts) + (1 if claim.assumptions else 0) + (1 if claim.reporting_period else 0)
    return min(CLAIM_BUDGET, 1 + len(claim.text) // 170 + extra)


def _paginate_claims(claims: Sequence[ExportClaim]) -> list[list[ExportClaim]]:
    pages: list[list[ExportClaim]] = [[]]
    used = 0
    for claim in claims:
        weight = _weight(claim)
        if pages[-1] and used + weight > CLAIM_BUDGET:
            pages.append([])
            used = 0
        pages[-1].append(claim)
        used += weight
    return pages if pages[-1] or len(pages) > 1 else [[]]


def _chunk[T](items: list[T], size: int) -> list[list[T]]:
    return [items[index : index + size] for index in range(0, len(items), size)]


def render_pptx(document: ExportDocument, theme: Theme) -> bytes:
    """Render one document in one theme."""
    deck = _Deck(document, theme)
    deck.cover()
    deck.changes()
    for section in document.sections:
        deck.section(section.title, section.claims, section.charts)
    if document.references:
        deck.sources(document.references)
    return deck.bytes()
