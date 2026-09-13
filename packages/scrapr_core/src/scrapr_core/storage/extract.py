"""Extracting and chunking uploaded documents (`REQ-DOC-001`, `-004`, `DEC-14`).

Implementation plan §10 sets the shape: extract per format, then chunk **with a
locator** so a citation resolves to a place in the file rather than to the file
as a whole. A citation that says "somewhere in this 80-page PDF" is not a
citation.

**Type comes from the bytes, not the filename** (`DEC-14`). A `.pdf` extension
is a claim the client makes; the magic bytes are what the file is. `Upload`'s
own docstring already says the stored content type is "what the file actually
is, not what the client asserted", and this is where that becomes true.

**Extraction failure is a state, never an empty document** (`AC-3`). A scanned
PDF holds no extractable text, and the honest outcome is a `failed` upload
saying so — not a `ready` one with nothing in it, which would look to every
later stage like a document the user uploaded and that said nothing.

**Everything extracted is untrusted** (`REQ-DOC-009`, §9). A document is as
capable of carrying an injection payload as a web page, and the fact that the
reader chose this file changes nothing: they may have been sent it.
"""

from __future__ import annotations

import csv
import io
import logging
import re
import zipfile
from collections.abc import Iterator, Sequence
from dataclasses import dataclass
from typing import Final, final

from scrapr_core.domain.json import JsonMapping

__all__ = [
    "ACCEPTED_TYPES",
    "CHUNK_CHARS",
    "ExtractedChunk",
    "ExtractionError",
    "detect_content_type",
    "extract_chunks",
]

# `DEC-14`. The message `REQ-DOC-001 AC-2` requires is generated from this, so
# the list and what a rejection says can never disagree.
ACCEPTED_TYPES: Final[dict[str, str]] = {
    "application/pdf": "PDF",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "Word (.docx)",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "Excel (.xlsx)",
    "text/csv": "CSV",
    "text/plain": "plain text",
    "text/markdown": "Markdown",
}

CHUNK_CHARS: Final = 1200
"""Roughly a long paragraph.

Big enough that a chunk usually contains the sentence *and* what qualifies it,
small enough that a citation points somewhere a reader can find. `REQ-DOC-005
AC-2` wants passages sufficient to support a claim, and a passage that needs
its neighbour to make sense is not one.
"""

logger = logging.getLogger(__name__)

_WHITESPACE = re.compile(r"[ \t]+")

_BLANK_LINES = re.compile(r"\n{3,}")
"""Runs of blank lines, collapsed when tidying. Three or more, so an ordinary
paragraph break survives."""

_PARAGRAPH_BREAK = re.compile(r"\n[ \t]*\n")
"""Where a plain-text document changes subject.

Two newlines, not three. Splitting on `_BLANK_LINES` instead was a real defect:
a normally formatted `.md` or `.txt` file separates paragraphs with a single
blank line, so nothing ever split — the whole document became one chunk,
unbounded in size and cited as "line 1" however long it was.
"""


class ExtractionError(RuntimeError):
    """The file could not be read as the type it claims to be.

    Raised so the caller can mark the upload `failed` with a reason
    (`REQ-DOC-004 AC-3`). Never swallowed into an empty result: a document that
    could not be read and a document that said nothing look identical
    downstream, and only one of them is a problem to tell the user about.
    """


@final
@dataclass(frozen=True, slots=True)
class ExtractedChunk:
    """One span of a document, and where in the file it came from."""

    ordinal: int
    text: str
    locator: JsonMapping
    """Page, sheet, or line range — whatever addresses this span in the
    original. Stored as JSON because the shape differs per format and a column
    per format would be a schema that grows with `DEC-14`."""


def detect_content_type(data: bytes, filename: str) -> str | None:
    """What this file actually is, or `None` if it is not a type we accept.

    Magic bytes first, extension only for the plain-text formats that have no
    signature to check. A `.csv` and a `.txt` are both text and telling them
    apart matters only for how they are chunked, so the filename is allowed to
    decide that and nothing else.
    """
    if data.startswith(b"%PDF-"):
        return "application/pdf"

    # Both OOXML formats are zip archives; the difference is what is inside.
    # Read from the archive's own directory, not by searching the raw bytes: a
    # member's name sits wherever earlier members' compressed sizes put it —
    # which moves with the timestamps inside them — and three bytes such as
    # `xl/` turn up inside compressed data by chance. A byte search therefore
    # misread the same kind of file on some saves and not others.
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = set(archive.namelist())
        except (zipfile.BadZipFile, ValueError):
            return None
        if "word/document.xml" in names:
            return (
                "application/vnd.openxmlformats-officedocument"
                ".wordprocessingml.document"
            )
        if "xl/workbook.xml" in names:
            return (
                "application/vnd.openxmlformats-officedocument"
                ".spreadsheetml.sheet"
            )
        return None

    # Text has no signature. Decodability is the test, and the extension picks
    # between the text formats.
    try:
        data[:8192].decode("utf-8")
    except UnicodeDecodeError:
        return None

    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == "csv":
        return "text/csv"
    if suffix in {"md", "markdown"}:
        return "text/markdown"
    if suffix in {"txt", ""}:
        return "text/plain"
    return None


def extract_chunks(data: bytes, content_type: str) -> list[ExtractedChunk]:
    """Text and locators from one document.

    Raises `ExtractionError` when the file cannot be read, and returns an empty
    list only when it genuinely holds no text — a blank page. The caller treats
    both as `failed`, because a document with nothing in it is not evidence
    either, but the distinction is kept here so the error message can say which
    happened.
    """
    if content_type == "application/pdf":
        return _chunk(_pdf_pages(data))
    if content_type.endswith("wordprocessingml.document"):
        return _chunk(_docx_paragraphs(data))
    if content_type.endswith("spreadsheetml.sheet"):
        return _chunk(_xlsx_rows(data))
    if content_type == "text/csv":
        return _chunk(_csv_rows(data))
    if content_type in {"text/plain", "text/markdown"}:
        return _chunk(_text_lines(data))

    raise ExtractionError(f"no extractor for {content_type}")


# --------------------------------------------------------------------------
# Per format: yield (text, locator) pairs
# --------------------------------------------------------------------------


def _pdf_pages(data: bytes) -> Iterator[tuple[str, JsonMapping]]:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = list(reader.pages)
    except (PyPdfError, ValueError, OSError) as exc:
        raise ExtractionError(f"this PDF could not be read: {exc}") from exc

    for number, page in enumerate(pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            # Skipping the page is right: the other forty are still evidence,
            # and failing the whole upload over one malformed content stream
            # would throw away a document the reader can see is fine. Logged,
            # never silent — a page that vanishes without a trace is content
            # the reader supplied and we quietly lost. The page number is safe
            # to log; the text is not.
            logger.warning(
                "skipped page %d of a PDF: %s", number, type(exc).__name__
            )
            continue
        if text.strip():
            yield text, {"page": number}


def _docx_paragraphs(data: bytes) -> Iterator[tuple[str, JsonMapping]]:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(data))
    except (
        # `BadZipFile` is what a truncated or corrupt `.docx` actually raises,
        # and it is not an `OSError`. Without it here the worker's extraction
        # step died on a bad file instead of marking the upload failed with a
        # reason, which is `REQ-DOC-004 AC-3` exactly.
        zipfile.BadZipFile,
        PackageNotFoundError,
        ValueError,
        KeyError,
        OSError,
    ) as exc:
        raise ExtractionError(f"this Word document could not be read: {exc}") from exc

    for index, paragraph in enumerate(document.paragraphs, start=1):
        if paragraph.text.strip():
            yield paragraph.text, {"paragraph": index}

    # Tables carry the numbers in most filings and reports, and skipping them
    # would drop exactly the evidence this feature exists to find.
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                yield " | ".join(cells), {"table": table_index, "row": row_index}


def _xlsx_rows(data: bytes) -> Iterator[tuple[str, JsonMapping]]:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        # `data_only`: the cached value of a formula, not the formula. A reader
        # checking a citation wants the number the sheet showed.
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (
        zipfile.BadZipFile,
        InvalidFileException,
        KeyError,
        ValueError,
        OSError,
    ) as exc:
        raise ExtractionError(f"this spreadsheet could not be read: {exc}") from exc

    for sheet in workbook.worksheets:
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = [str(cell) for cell in row if cell is not None]
            if cells:
                yield " | ".join(cells), {"sheet": sheet.title, "row": row_index}
    workbook.close()


def _csv_rows(data: bytes) -> Iterator[tuple[str, JsonMapping]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ExtractionError("this CSV is not valid UTF-8 text") from exc

    for index, row in enumerate(csv.reader(io.StringIO(text)), start=1):
        cells = [cell.strip() for cell in row if cell.strip()]
        if cells:
            yield " | ".join(cells), {"row": index}


def _text_lines(data: bytes) -> Iterator[tuple[str, JsonMapping]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ExtractionError("this file is not valid UTF-8 text") from exc

    line = 1
    for block in _PARAGRAPH_BREAK.split(text):
        cleaned = block.strip()
        if cleaned:
            yield cleaned, {"line": line}
        # Two for the break itself, plus however many the block spanned.
        line += block.count("\n") + 2


# --------------------------------------------------------------------------


def _chunk(pieces: Iterator[tuple[str, JsonMapping]]) -> list[ExtractedChunk]:
    """Group extracted pieces into chunks, keeping the first locator.

    Pieces are accumulated until `CHUNK_CHARS`, and the chunk is addressed by
    where it *starts*. A chunk spanning pages 3 and 4 is cited as page 3, which
    is where a reader opening the file should begin looking — precise enough to
    find, and honest about being approximate.
    """
    chunks: list[ExtractedChunk] = []
    buffer: list[str] = []
    start: JsonMapping | None = None
    length = 0

    def flush() -> None:
        nonlocal buffer, start, length
        if buffer and start is not None:
            chunks.append(
                ExtractedChunk(
                    ordinal=len(chunks),
                    text=_tidy("\n".join(buffer)),
                    locator=start,
                )
            )
        buffer, start, length = [], None, 0

    for text, locator in pieces:
        # One piece can be longer than a whole chunk — a dense PDF page, a
        # single unbroken paragraph. Without this it would become one chunk of
        # whatever size it happened to be, and the citation would point at a
        # wall of text rather than at a passage.
        for part in _split_long(text):
            if start is None:
                start = locator
            buffer.append(part)
            length += len(part)
            if length >= CHUNK_CHARS:
                flush()

    flush()
    return [chunk for chunk in chunks if chunk.text.strip()]


def _split_long(text: str) -> Iterator[str]:
    """Break an oversized piece on a sentence end, or on whitespace if it has none.

    Sentence ends first, because a passage cut mid-sentence cannot be quoted as
    an excerpt — and the excerpt is what `REQ-EVID-007` shows the reader.
    """
    if len(text) <= CHUNK_CHARS:
        yield text
        return

    remaining = text
    while len(remaining) > CHUNK_CHARS:
        window = remaining[:CHUNK_CHARS]
        cut = max(window.rfind(". "), window.rfind("\n"))
        if cut < CHUNK_CHARS // 2:
            cut = window.rfind(" ")
        if cut <= 0:
            # No break anywhere: one enormous token. Cut it rather than emit it
            # whole, because the alternative is no bound at all.
            cut = CHUNK_CHARS - 1
        yield remaining[: cut + 1].strip()
        remaining = remaining[cut + 1 :]

    if remaining.strip():
        yield remaining.strip()


def _tidy(text: str) -> str:
    """Collapse the whitespace PDF extraction leaves behind.

    Non-destructive in the sense that matters: no word is changed, so an
    excerpt taken from a chunk still appears in the document a reader opens.
    """
    return _BLANK_LINES.sub("\n\n", _WHITESPACE.sub(" ", text)).strip()


def accepted_types_message() -> str:
    """What `REQ-DOC-001 AC-2` requires a rejection to say."""
    return ", ".join(sorted(set(ACCEPTED_TYPES.values())))


def describe_locator(locator: Sequence[tuple[str, object]] | JsonMapping) -> str:
    """A locator in words, for a citation (`REQ-DOC-008 AC-2`)."""
    items = dict(locator)
    if "page" in items:
        return f"page {items['page']}"
    if "sheet" in items:
        return f"sheet {items['sheet']}, row {items.get('row', '?')}"
    if "table" in items:
        return f"table {items['table']}, row {items.get('row', '?')}"
    if "paragraph" in items:
        return f"paragraph {items['paragraph']}"
    if "row" in items:
        return f"row {items['row']}"
    if "line" in items:
        return f"line {items['line']}"
    return "the document"
