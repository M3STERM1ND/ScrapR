"""The two renderers and the six themes (`REQ-EXP-001..004`, `REQ-EXP-006`, `REQ-EXP-009`, `DEC-21`, `DEC-22`).

Pure: a synthetic `ExportDocument` in, bytes out, read back with pypdf and
python-pptx. The document deliberately carries the hard cases — every claim
type, a conflict, a forecast with assumptions, a line chart, a table, a
document source, a hostile excerpt and a `javascript:` URL — because a renderer
that handles the easy report proves nothing about the real ones.
"""

from __future__ import annotations

import datetime as dt
import io
import zipfile
from decimal import Decimal
from pathlib import Path

import pytest
from pptx import Presentation
from pypdf import PdfReader

from scrapr_core.db.enums import ClaimType, ExportTheme, VersionStatus, VizKind
from scrapr_core.export import THEMES, render_pdf, render_pptx
from scrapr_core.export.document import (
    ExportChange,
    ExportChart,
    ExportClaim,
    ExportConflict,
    ExportDocument,
    ExportReference,
    ExportSection,
    ExportSeries,
    clean_text,
    safe_link,
)
from scrapr_core.export.themes import contrast_ratio

HOSTILE = '<link href="javascript:alert(1)">click</link> & <b>bold</b>'
OBJECTIVE = "How is Acme Corp positioned against its competitors?"


def document(status: VersionStatus = VersionStatus.COMPLETE) -> ExportDocument:
    return ExportDocument(
        objective=OBJECTIVE,
        subject="Acme Corp",
        version_number=2,
        version_created_at=dt.datetime(2026, 9, 13, 9, 30, tzinfo=dt.UTC),
        status=status,
        changes_headline="1 meaningful change since version 1.",
        changes=(ExportChange(label="Figure changed", summary="Moved from $1.2bn to $1.5bn."),),
        sections=(
            ExportSection(
                title="Executive summary",
                is_executive_summary=True,
                claims=(
                    ExportClaim(
                        text="Acme reported revenue of $1.5bn for fiscal 2025.",
                        claim_type=ClaimType.FACT,
                        confidence="Moderate confidence",
                        references=(1, 2),
                        reporting_period="2025-12-31",
                        conflicts=(
                            ExportConflict(
                                resolved=False,
                                explanation="Nothing in the evidence accounts for the difference.",
                                sides=(),
                            ),
                        ),
                    ),
                    ExportClaim(
                        text="Acme's enterprise position looks strong.",
                        claim_type=ClaimType.ANALYSIS,
                        confidence="High confidence",
                        references=(1,),
                    ),
                ),
            ),
            ExportSection(
                title="Outlook",
                is_executive_summary=False,
                claims=(
                    ExportClaim(
                        text="Revenue is likely to keep growing next year.",
                        claim_type=ClaimType.FORECAST,
                        confidence="Low confidence",
                        references=(2,),
                        assumptions=("enterprise renewals continue",),
                    ),
                    ExportClaim(
                        text=f"Hiring plans could not be established. {HOSTILE}",
                        claim_type=ClaimType.UNCERTAINTY,
                        confidence=None,
                        references=(),
                    ),
                ),
                charts=(
                    ExportChart(
                        kind=VizKind.LINE,
                        title="Revenue by year",
                        unit="USD",
                        series=(
                            ExportSeries(
                                name="Revenue",
                                points=(
                                    ("2023", Decimal("900000000")),
                                    ("2024", Decimal("1200000000")),
                                    ("2025", Decimal("1500000000")),
                                ),
                            ),
                            ExportSeries(
                                name="Operating cost",
                                points=(("2024", Decimal("800000000")), ("2025", Decimal("950000000"))),
                            ),
                        ),
                        references=(1,),
                    ),
                    ExportChart(
                        kind=VizKind.TABLE,
                        title="Headcount",
                        unit=None,
                        series=(ExportSeries(name="Engineers", points=(("2025", Decimal("420")),)),),
                    ),
                ),
            ),
        ),
        references=(
            ExportReference(
                number=1, name="Reuters", publisher="Reuters", url="https://reuters.com/acme",
                tier="Established source", retrieved_at=dt.datetime(2026, 9, 12, tzinfo=dt.UTC),
            ),
            ExportReference(
                number=2, name="board-notes.pdf", publisher=None,
                url=safe_link("javascript:alert(document.cookie)"),
                tier="Established source", retrieved_at=dt.datetime(2026, 9, 12, tzinfo=dt.UTC),
                from_your_document=True, document_removed=True,
            ),
        ),
    )


def pdf_text(data: bytes) -> str:
    """Every page's text, whitespace collapsed, so a title a large theme wraps
    across lines still reads as the sentence it is."""
    raw = " ".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    return " ".join(raw.split())


def pptx_texts(data: bytes) -> list[str]:
    deck = Presentation(io.BytesIO(data))
    texts = []
    for slide in deck.slides:
        parts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        texts.append("\n".join(parts))
    return texts


# --------------------------------------------------------------------------
# Themes — `DEC-22`
# --------------------------------------------------------------------------


def test_all_six_themes_are_defined() -> None:
    """`REQ-EXP-003 AC-1`."""
    assert set(THEMES) == set(ExportTheme)


@pytest.mark.parametrize("key", list(ExportTheme))
def test_every_theme_meets_aa_contrast(key: ExportTheme) -> None:
    theme = THEMES[key]
    assert contrast_ratio(theme.text, theme.ground) >= 4.5
    assert contrast_ratio(theme.muted, theme.ground) >= 4.5
    # The accent carries citation markers, so it is text too.
    assert contrast_ratio(theme.accent, theme.ground) >= 4.5
    assert contrast_ratio(theme.accent_on_accent, theme.accent) >= 4.5


def test_the_themes_are_visually_distinct() -> None:
    """`REQ-EXP-003 AC-2`."""
    signatures = {
        (theme.ground, theme.accent, theme.heading_serif, theme.body_serif, theme.cover)
        for theme in THEMES.values()
    }
    assert len(signatures) == len(THEMES)


# --------------------------------------------------------------------------
# PDF — `REQ-EXP-001`
# --------------------------------------------------------------------------


@pytest.mark.parametrize("key", list(ExportTheme))
def test_the_pdf_carries_the_whole_report_in_every_theme(key: ExportTheme) -> None:
    """`REQ-EXP-001 AC-1`, `REQ-EXP-006 AC-2`, `REQ-EXP-009`."""
    data = render_pdf(document(), THEMES[key])

    assert data.startswith(b"%PDF-")
    text = pdf_text(data)
    for expected in (
        "How is Acme Corp positioned",
        "Version 2",
        "Sep 13, 2026",
        "Executive summary",
        "Fact",
        "Analysis",
        "Forecast",
        "Uncertain",
        "Moderate confidence",
        "[1][2]",
        "Sources disagree",
        "Assumptions: enterprise renewals continue",
        "Revenue by year",
        "Headcount",
        "Reuters",
        "your document",
        "file since deleted",
        "What",  # What's changed
        "Moved from $1.2bn to $1.5bn",
    ):
        assert expected in text, (key, expected)


def test_the_pdf_prints_hostile_text_and_never_links_it() -> None:
    """`REQ-SEC-013`: retrieved markup is content, not structure."""
    data = render_pdf(document(), THEMES[ExportTheme.PROFESSIONAL])

    assert "<b>bold</b>" in pdf_text(data)
    assert b"javascript" not in data.lower()
    assert b"https://reuters.com/acme" in data


def test_the_pdf_is_deterministic_and_every_theme_differs() -> None:
    """Same version and theme, same bytes (`NFR-REL-003`); different theme,
    different document (`REQ-EXP-003 AC-2`)."""
    first = render_pdf(document(), THEMES[ExportTheme.INVESTOR])
    assert render_pdf(document(), THEMES[ExportTheme.INVESTOR]) == first

    outputs = {render_pdf(document(), theme) for theme in THEMES.values()}
    assert len(outputs) == len(THEMES)


def test_a_partial_version_says_so_on_the_cover() -> None:
    text = pdf_text(render_pdf(document(VersionStatus.PARTIAL), THEMES[ExportTheme.MINIMAL]))
    assert "could not be answered" in text


def test_the_pdf_embeds_its_fonts() -> None:
    data = render_pdf(document(), THEMES[ExportTheme.MINIMAL])
    assert b"/FontFile2" in data


# --------------------------------------------------------------------------
# PowerPoint — `REQ-EXP-002`
# --------------------------------------------------------------------------


@pytest.mark.parametrize("key", list(ExportTheme))
def test_the_deck_distributes_the_report_across_slides(key: ExportTheme) -> None:
    """`REQ-EXP-002 AC-1`, `REQ-EXP-006 AC-2`."""
    data = render_pptx(document(), THEMES[key])
    texts = pptx_texts(data)
    joined = "\n".join(texts)

    # Cover, what's changed, two section slides, two chart slides, sources.
    assert len(texts) >= 7
    assert all("Version 2 · generated from research dated Sep 13, 2026" in text for text in texts)
    for expected in ("Fact · Moderate confidence", "Analysis", "Forecast", "Uncertain", "[1]",
                     "Sources disagree", "Revenue by year", "Reuters", "your document"):
        assert expected in joined, (key, expected)


def test_charts_are_native_and_tables_are_tables() -> None:
    """`REQ-EXP-002 AC-2`: visualizations are slide content, not placeholders."""
    deck = Presentation(io.BytesIO(render_pptx(document(), THEMES[ExportTheme.CORPORATE])))
    charts = [shape for slide in deck.slides for shape in slide.shapes if shape.has_chart]
    tables = [shape for slide in deck.slides for shape in slide.shapes if shape.has_table]

    assert len(charts) == 1
    chart = charts[0].chart
    assert [series.name for series in chart.plots[0].series] == ["Revenue", "Operating cost"]
    assert list(chart.plots[0].categories) == ["2023", "2024", "2025"]
    assert len(tables) == 1


def test_the_deck_is_a_well_formed_package() -> None:
    """`REQ-EXP-002 AC-3`: every part parses, and nothing invalid reaches XML."""
    data = render_pptx(document(), THEMES[ExportTheme.DARK])
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        assert package.testzip() is None
        assert "[Content_Types].xml" in package.namelist()
        from xml.dom.minidom import parseString

        for name in package.namelist():
            if name.endswith(".xml") or name.endswith(".rels"):
                # The deck this test just rendered, not outside input.
                parseString(package.read(name))  # noqa: S318 - raises on malformed XML
    # And python-pptx itself reads it back.
    assert len(Presentation(io.BytesIO(data)).slides) > 0


def test_the_deck_never_links_an_unsafe_url() -> None:
    data = render_pptx(document(), THEMES[ExportTheme.PROFESSIONAL])
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        rels = b"".join(package.read(name) for name in package.namelist() if name.endswith(".rels"))
    assert b"javascript" not in rels.lower()
    assert b"https://reuters.com/acme" in rels


@pytest.mark.parametrize("key", list(ExportTheme))
def test_every_slide_uses_its_themes_ground(key: ExportTheme) -> None:
    theme = THEMES[key]
    deck = Presentation(io.BytesIO(render_pptx(document(), theme)))
    grounds = {str(slide.background.fill.fore_color.rgb) for slide in deck.slides}
    assert grounds == {theme.ground.lstrip("#").upper()}


def test_every_theme_renders_a_different_deck() -> None:
    """`REQ-EXP-003 AC-2`, in the second format. Themes sharing a white ground
    still differ in accent, type and cover."""
    def slides_xml(data: bytes) -> bytes:
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            return b"".join(
                package.read(name)
                for name in sorted(package.namelist())
                if name.startswith("ppt/slides/slide")
            )

    decks = {slides_xml(render_pptx(document(), theme)) for theme in THEMES.values()}
    assert len(decks) == len(THEMES)


# --------------------------------------------------------------------------
# Boundaries
# --------------------------------------------------------------------------


def test_text_is_cleaned_of_characters_xml_cannot_hold() -> None:
    assert clean_text("Acme\x00 reported\x0b  revenue\n of $1bn") == "Acme reported revenue of $1bn"


@pytest.mark.parametrize(
    ("url", "linked"),
    [
        ("https://reuters.com/acme", True),
        ("http://example.com/a", True),
        ("javascript:alert(1)", False),
        ("file:///etc/passwd", False),
        ("data:text/html,hi", False),
        ("https://example.com/with space", False),
        (None, False),
    ],
)
def test_only_web_links_are_made_clickable(url: str | None, linked: bool) -> None:
    assert (safe_link(url) is not None) is linked


def test_the_export_package_cannot_reach_a_model() -> None:
    """`REQ-EXP-004 AC-1`, `REQ-EXP-005 AC-2`: no model call at export time."""
    package = Path(__file__).resolve().parents[2] / "src" / "scrapr_core" / "export"
    for source in package.glob("*.py"):
        text = source.read_text(encoding="utf-8")
        assert "scrapr_core.llm" not in text, source.name
        assert "anthropic" not in text.lower(), source.name
